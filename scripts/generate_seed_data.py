"""Generate the deterministic Firestore product seed from Data.xlsx."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Data.xlsx"
OUTPUT = ROOT / "data" / "products.json"


def number(value):
    if value in (None, ""):
        return None
    return round(float(value), 2)


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


workbook = load_workbook(SOURCE, data_only=True, read_only=False)
sheet = workbook["Sheet1"]
products = []

for row in sheet.iter_rows(min_row=2, max_col=9, values_only=True):
    sku, name, category, stock, price, cost, pack_qty, pieces_per_pack, pack_price = row
    if not sku or not name:
        continue
    price_value = number(price)
    products.append(
        {
            "id": str(sku).strip(),
            "sku": str(sku).strip(),
            "name": str(name).strip(),
            "description": "",
            "category": str(category).strip(),
            "categoryId": slug(str(category)),
            "stock": int(stock or 0),
            "price": price_value,
            "costPrice": number(cost),
            "wholesalePackQuantity": number(pack_qty),
            "piecesPerPack": number(pieces_per_pack),
            "wholesalePackPrice": number(pack_price),
            "active": True,
            "archived": False,
            "sellable": bool(price_value and price_value > 0),
            "pinned": False,
            "lowStockLevel": 8,
            "imageUrl": "",
            "source": "Data.xlsx",
        }
    )

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(products, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(f"Wrote {len(products)} products to {OUTPUT}")

